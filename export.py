"""Excel output of a reconciliation run.

Sheets
  Summary           run details, total amounts (net sales, net settlement, difference, matched, mismatched),
                    amounts by status, exception counts by type, hygiene check overview, definitions
  Hygiene Checks    EVERY hygiene check per dataset with PASS / FAIL / NOT RUN - including checks with no findings
  Results           one row per key: all amounts + matched / mismatched amount + hygiene flags, with a TOTAL row
  Exceptions        reconciliation exceptions, row level
  Hygiene Findings  hygiene findings, row level, labelled with the readable check name
  Column Mapping    which file column was used for which field in this run

Written with openpyxl in write-only (streaming) mode so large runs do not have to be held in memory twice.
"""
import io
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import db
import hygiene
import mapping as colmap
import recon
from schema import DATASETS

MONEY = "#,##0.00;[Red]-#,##0.00"
COUNT = "#,##0"
MAX_ROWS = 1_000_000            # Excel holds 1,048,576 rows per sheet; longer tables continue on a second sheet

STATUS_LABELS = {
    "MATCHED": "Matched",
    "UNMATCHED": "Unmatched (amount differs)",
    "SALES_NOT_IN_SETTLEMENT": "Sales not in settlement",
    "SETTLEMENT_NOT_IN_SALES": "Settlement not in sales",
}
STATUS_ORDER = list(STATUS_LABELS)

# Single definition of the two derived amounts, used by the screens and the Excel output
MATCHED_AMOUNT_SQL = "CASE WHEN status = 'MATCHED' THEN COALESCE(net_sales, 0) ELSE 0 END"
MISMATCHED_AMOUNT_SQL = "CASE WHEN status = 'MATCHED' THEN 0 ELSE abs(COALESCE(difference, 0)) END"

DEFINITIONS = [
    ("Net sales", "Forward sales amount minus the return amount (RT if the packet is in RT, otherwise RTO), per sales key."),
    ("Net settlement", "Forward settlement amount minus the sum of reverse settlement amounts, per settlement key."),
    ("Difference", "Net sales minus net settlement. Positive = settled short, negative = settled in excess."),
    ("Matched amount", "Net sales of the keys whose status is Matched (|difference| within the tolerance). 0 for every other key."),
    ("Mismatched amount", "|Difference| of every key that is not Matched: the short / excess amount for Unmatched keys, "
                          "the full net sales for 'Sales not in settlement' and the full net settlement for 'Settlement not in sales'."),
]

_HEAD_FILL = PatternFill("solid", fgColor="00338D")
_HEAD_FONT = Font(bold=True, color="FFFFFF")
_SECTION_FONT = Font(bold=True, size=12, color="00338D")
_BOLD = Font(bold=True)
_TOTAL_FILL = PatternFill("solid", fgColor="EEF2F7")
_RESULT_FILL = {"PASS": PatternFill("solid", fgColor="E0F2E6"), "FAIL": PatternFill("solid", fgColor="FDE3DF"),
                "NOT RUN": PatternFill("solid", fgColor="E5E9EF")}
_WRAP = Alignment(wrap_text=True, vertical="top")


def label(code):
    """Readable name of a status / hygiene check / reconciliation exception code."""
    if code in STATUS_LABELS:
        return STATUS_LABELS[code]
    if code in hygiene.CHECKS:
        return hygiene.CHECKS[code][0]
    if code in recon.EXCEPTIONS:
        return recon.EXCEPTIONS[code][0]
    return str(code).replace("_", " ").capitalize()


def dataset_label(name):
    return DATASETS[name]["label"] if name in DATASETS else (name or "Net sales vs net settlement")


def _num(v):
    if isinstance(v, Decimal):
        return float(v)
    if isinstance(v, str):
        # control characters from the uploaded files make openpyxl refuse to write the workbook
        return v if v.isprintable() else ILLEGAL_CHARACTERS_RE.sub("", v)   # fast path for normal text
    return v


def _cell(ws, value, font=None, fill=None, fmt=None, wrap=False):
    c = WriteOnlyCell(ws, value=_num(value))
    if font:
        c.font = font
    if fill:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    if wrap:
        c.alignment = _WRAP
    return c


def _sheet(wb, title, widths, freeze="A2"):
    ws = wb.create_sheet(title)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    if freeze:
        ws.freeze_panes = freeze
    return ws


def _header(ws, headers):
    ws.append([_cell(ws, h, font=_HEAD_FONT, fill=_HEAD_FILL) for h in headers])


def _table(wb, title, headers, widths, rows, formats=None, total=None):
    """Write a row-level table. rows = iterable of lists; formats = {column index: number format}.
    total = list for a closing TOTAL row (None cells stay empty). Continues on '<title> (2)' past MAX_ROWS."""
    formats = formats or {}
    part, n, ws = 1, 0, None

    def new_sheet():
        nonlocal ws, n
        ws = _sheet(wb, title if part == 1 else f"{title} ({part})", widths)
        _header(ws, headers)
        n = 0

    new_sheet()
    for r in rows:
        if n >= MAX_ROWS:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{n + 1}"
            part += 1
            new_sheet()
        ws.append([_cell(ws, v, fmt=formats[i]) if i in formats and v is not None else _num(v)
                   for i, v in enumerate(r)])
        n += 1
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(n, 1) + 1}"
    if total is not None:
        ws.append([_cell(ws, v, font=_BOLD, fill=_TOTAL_FILL, fmt=formats.get(i)) for i, v in enumerate(total)])
    return ws


def amount_totals(cur, run_id):
    """Amounts by status + grand total for a run (also used by the Summary screen)."""
    cur.execute(f"""SELECT status, count(*) AS n,
                           COALESCE(sum(net_sales), 0)      AS net_sales,
                           COALESCE(sum(net_settlement), 0) AS net_settlement,
                           COALESCE(sum(difference), 0)     AS difference,
                           COALESCE(sum({MATCHED_AMOUNT_SQL}), 0)    AS matched_amount,
                           COALESCE(sum({MISMATCHED_AMOUNT_SQL}), 0) AS mismatched_amount
                    FROM recon_results WHERE run_id = %s GROUP BY status""", (run_id,))
    by_status = {r["status"]: r for r in cur.fetchall()}
    rows = [by_status[s] for s in STATUS_ORDER if s in by_status] + \
           [r for s, r in by_status.items() if s not in STATUS_ORDER]
    keys = ("n", "net_sales", "net_settlement", "difference", "matched_amount", "mismatched_amount")
    grand = {k: sum((r[k] for r in rows), Decimal(0)) for k in keys}
    grand["n"] = int(grand["n"])
    return rows, grand


def check_log(cur, run_id, category):
    cur.execute("SELECT * FROM recon_check_log WHERE run_id = %s AND category = %s ORDER BY seq", (run_id, category))
    return cur.fetchall()


def _summary_sheet(wb, run, totals, grand, recon_log, exc_counts, hyg_log, hyg_counts):
    ws = _sheet(wb, "Summary", [46, 22, 22, 22, 22, 22, 22], freeze=None)

    def section(title, headers=None):
        ws.append([])
        ws.append([_cell(ws, title, font=_SECTION_FONT)])
        if headers:
            _header(ws, headers)

    ws.append([_cell(ws, "AR Reconciliation - Sales vs Settlement", font=Font(bold=True, size=14, color="00338D"))])
    ws.append([f"Run #{run['run_id']} - {run['recon_month']:%B %Y}"])

    section("Run details", ["Field", "Value"])
    for k, v in [("Run", run["run_id"]), ("Run time", f"{run['created_at']:%d %b %Y %H:%M}"),
                 ("Reconciliation month", f"{run['recon_month']:%B %Y}"),
                 ("Amount tolerance (INR)", run["tolerance"]), ("Return window (days)", run["return_window"]),
                 ("Commission / tax tolerance (INR)", run["commission_tol"]),
                 ("Platform filter", run["platform_filter"] or "(all platforms)")]:
        ws.append([k, _num(v)])

    section("Amounts", ["Measure", "Amount (INR)"])
    for name, key in [("Net sales (total)", "net_sales"), ("Net settlement (total)", "net_settlement"),
                      ("Difference (net sales - net settlement)", "difference"),
                      ("Matched amount", "matched_amount"), ("Mismatched amount", "mismatched_amount")]:
        ws.append([_cell(ws, name, font=_BOLD), _cell(ws, grand[key], font=_BOLD, fmt=MONEY)])

    section("Amounts by status", ["Status", "Keys", "Net sales", "Net settlement",
                                  "Difference (net sales - net settlement)", "Matched amount", "Mismatched amount"])
    cols = ("net_sales", "net_settlement", "difference", "matched_amount", "mismatched_amount")
    for t in totals:
        ws.append([label(t["status"]), _cell(ws, t["n"], fmt=COUNT)] + [_cell(ws, t[c], fmt=MONEY) for c in cols])
    ws.append([_cell(ws, "TOTAL", font=_BOLD, fill=_TOTAL_FILL), _cell(ws, grand["n"], font=_BOLD, fill=_TOTAL_FILL, fmt=COUNT)] +
              [_cell(ws, grand[c], font=_BOLD, fill=_TOTAL_FILL, fmt=MONEY) for c in cols])

    section("Counts", ["Measure", "Count"])
    for k, v in [("Sales keys (net_sales rows)", run["total_sales"]),
                 ("Settlement keys (net_settlement rows)", run["total_settlement"]),
                 ("Matched", run["matched"]), ("Unmatched (amount differs)", run["unmatched"]),
                 ("Sales not in settlement", run["sales_only"]), ("Settlement not in sales", run["settlement_only"]),
                 ("Reconciliation exceptions (rows on sheet 'Exceptions')", run["exception_count"]),
                 ("Hygiene findings (rows on sheet 'Hygiene Findings')", run["hygiene_count"])]:
        ws.append([k, _cell(ws, v, fmt=COUNT)])

    section("Reconciliation exceptions by type", ["Exception", "Code", "Rows", "Result", "Note"])
    if recon_log:
        seen = {}
        for c in recon_log:          # one line per exception type (some run once per settlement file)
            s = seen.setdefault(c["check_code"], {"name": c["check_name"], "ran": False, "note": set()})
            s["ran"] = s["ran"] or c["status"] != "NOT RUN"
            if c["status"] == "NOT RUN" and c["note"]:
                s["note"].add(f"{dataset_label(c['dataset'])}: {c['note']}")
        for code, s in seen.items():
            n = exc_counts.get(code, 0)
            result = "FAIL" if n else ("PASS" if s["ran"] else "NOT RUN")
            ws.append([s["name"], code, _cell(ws, n, fmt=COUNT), _cell(ws, result, fill=_RESULT_FILL[result]),
                       "; ".join(sorted(s["note"]))])
    else:
        for code, n in sorted(exc_counts.items()):
            ws.append([label(code), code, _cell(ws, n, fmt=COUNT)])

    section("Hygiene checks (detail per dataset on sheet 'Hygiene Checks')",
            ["Hygiene check", "Code", "Datasets checked", "Datasets failed", "Findings (rows)", "Result"])
    if hyg_log:
        agg = {}
        for c in hyg_log:
            a = agg.setdefault(c["check_code"], {"name": c["check_name"], "ran": 0, "failed": 0})
            a["ran"] += c["status"] != "NOT RUN"
            a["failed"] += c["status"] == "FAIL"
        for code, a in agg.items():
            result = "FAIL" if a["failed"] else ("PASS" if a["ran"] else "NOT RUN")
            ws.append([a["name"], code, a["ran"], a["failed"], _cell(ws, hyg_counts.get(code, 0), fmt=COUNT),
                       _cell(ws, result, fill=_RESULT_FILL[result])])
    else:
        for code, n in sorted(hyg_counts.items()):
            ws.append([label(code), code, None, None, _cell(ws, n, fmt=COUNT), "FAIL"])

    section("Definitions", ["Term", "Meaning"])
    for term, meaning in DEFINITIONS:
        ws.append([_cell(ws, term, font=_BOLD), meaning])


def _hygiene_checks_sheet(wb, hyg_log):
    ws = _sheet(wb, "Hygiene Checks", [5, 40, 30, 62, 30, 44, 16, 16, 11, 60])
    _header(ws, ["#", "Hygiene check", "Check code", "What it tests", "Dataset", "Column(s) checked",
                 "Records checked", "Records failed", "Result", "Note"])
    if not hyg_log:
        ws.append([None, "The list of checks was not recorded for this run (it was made before checks were logged). "
                         "Run the reconciliation again to get it."])
        return
    for i, c in enumerate(hyg_log, start=1):
        ws.append([i, c["check_name"], c["check_code"], c["description"], dataset_label(c["dataset"]),
                   c["columns_checked"],
                   _cell(ws, c["records_checked"], fmt=COUNT), _cell(ws, c["records_failed"], fmt=COUNT),
                   _cell(ws, c["status"], font=_BOLD, fill=_RESULT_FILL.get(c["status"])), c["note"]])
    ws.auto_filter.ref = f"A1:J{len(hyg_log) + 1}"


def build_workbook(run):
    """Return a BytesIO holding the .xlsx for the run (row of recon_runs)."""
    run_id = run["run_id"]
    wb = Workbook(write_only=True)

    with db.cursor(commit=False) as cur:
        totals, grand = amount_totals(cur, run_id)
        recon_log = check_log(cur, run_id, "RECONCILIATION")
        hyg_log = check_log(cur, run_id, "HYGIENE")
        cur.execute("""SELECT category, exception_type, count(*) AS n FROM recon_exceptions
                       WHERE run_id = %s GROUP BY 1, 2""", (run_id,))
        counts = cur.fetchall()
        exc_counts = {r["exception_type"]: r["n"] for r in counts if r["category"] == "RECONCILIATION"}
        hyg_counts = {r["exception_type"]: r["n"] for r in counts if r["category"] == "HYGIENE"}

        _summary_sheet(wb, run, totals, grand, recon_log, exc_counts, hyg_log, hyg_counts)
        _hygiene_checks_sheet(wb, hyg_log)

        # ---- Results: every key with all its amounts and the hygiene checks its rows failed
        cur.execute(f"""
            WITH hs AS (SELECT packet_id AS k, array_agg(DISTINCT exception_type) AS t FROM recon_exceptions
                        WHERE run_id = %s AND category = 'HYGIENE' AND dataset LIKE 'sales%%' AND packet_id IS NOT NULL
                        GROUP BY 1),
                 ht AS (SELECT order_release_id AS k, array_agg(DISTINCT exception_type) AS t FROM recon_exceptions
                        WHERE run_id = %s AND category = 'HYGIENE' AND dataset LIKE 'settlement%%' AND order_release_id IS NOT NULL
                        GROUP BY 1)
            SELECT r.status, r.packet_id, r.order_release_id, r.forward_sales, r.return_source, r.return_amount,
                   r.net_sales, r.forward_settlement, r.reverse_settlement, r.net_settlement, r.difference,
                   {MATCHED_AMOUNT_SQL} AS matched_amount, {MISMATCHED_AMOUNT_SQL} AS mismatched_amount,
                   hs.t AS sales_flags, ht.t AS settlement_flags
            FROM recon_results r
            LEFT JOIN hs ON hs.k = r.packet_id
            LEFT JOIN ht ON ht.k = r.order_release_id
            WHERE r.run_id = %s ORDER BY r.status, r.packet_id, r.order_release_id""", (run_id, run_id, run_id))
        check_order = {code: i for i, code in enumerate(hygiene.CHECKS)}

        def result_rows():
            for r in cur.fetchall():
                flags = sorted(set(r["sales_flags"] or []) | set(r["settlement_flags"] or []),
                               key=lambda c: check_order.get(c, 99))
                yield [label(r["status"]), r["packet_id"], r["order_release_id"], r["forward_sales"],
                       None if r["return_source"] in (None, "NONE") else r["return_source"], r["return_amount"],
                       r["net_sales"], r["forward_settlement"], r["reverse_settlement"], r["net_settlement"],
                       r["difference"], r["matched_amount"], r["mismatched_amount"],
                       "; ".join(label(c) for c in flags) or None]

        money_cols = {i: MONEY for i in (3, 5, 6, 7, 8, 9, 10, 11, 12)}
        _table(wb, "Results",
               ["Status", "Packet ID (sales)", "Order release ID (settlement)", "Forward sales", "Return type",
                "Return amount", "Net sales", "Forward settlement", "Reverse settlement", "Net settlement",
                "Difference (net sales - net settlement)", "Matched amount", "Mismatched amount", "Hygiene flags"],
               [28, 20, 26, 16, 12, 16, 16, 18, 18, 16, 24, 18, 18, 60],
               result_rows(), formats=money_cols,
               total=[f"TOTAL ({grand['n']:,} keys)", None, None, None, None, None, grand["net_sales"], None, None,
                      grand["net_settlement"], grand["difference"], grand["matched_amount"],
                      grand["mismatched_amount"], None])

        # ---- Row-level exceptions and hygiene findings, labelled with the readable names
        def exception_rows(category):
            cur.execute("""SELECT exception_type, dataset, row_no, packet_id, order_release_id,
                                  amount_1, amount_2, difference, description
                           FROM recon_exceptions WHERE run_id = %s AND category = %s
                           ORDER BY exception_type, dataset, row_no, packet_id, order_release_id""", (run_id, category))
            for e in cur.fetchall():
                yield [label(e["exception_type"]), e["exception_type"], dataset_label(e["dataset"]), e["row_no"],
                       e["packet_id"], e["order_release_id"], e["amount_1"], e["amount_2"], e["difference"],
                       e["description"]]

        widths = [42, 38, 30, 12, 20, 22, 20, 22, 18, 90]
        fm = {6: MONEY, 7: MONEY, 8: MONEY}
        _table(wb, "Exceptions",
               ["Exception", "Exception code", "Source", "File row no", "Packet ID", "Order release ID",
                "Net sales / expected", "Net settlement / actual", "Difference", "Detail"],
               widths, exception_rows("RECONCILIATION"), formats=fm)
        _table(wb, "Hygiene Findings",
               ["Hygiene check", "Check code", "Dataset", "File row no", "Packet ID", "Order release ID",
                "Sale / expected value", "Return / actual value", "Difference (amount or days)", "Detail"],
               widths, exception_rows("HYGIENE"), formats=fm)

    # ---- Column mapping used by the run
    m = colmap.Mapping.from_json(run.get("column_mapping"))
    ws = _sheet(wb, "Column Mapping", [14, 34, 34, 10, 34, 34])
    _header(ws, ["Side", "File", "Field used by the logic", "Required", "Column chosen (as in the file)", "Table column"])
    if m:
        for r in colmap.describe(m):
            ws.append([_num(r[k]) for k in ("side", "dataset", "field", "required", "file_column", "table_column")])
    else:
        ws.append([None, "Not recorded for this run (made before column mapping was introduced); "
                         "the standard column names were used."])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


RESULT_HEADERS = ["Status", "Packet ID (sales)", "Order release ID (settlement)", "Forward sales", "Return type",
                  "Return amount", "Net sales", "Forward settlement", "Reverse settlement", "Net settlement",
                  "Difference (net sales - net settlement)", "Matched amount", "Mismatched amount"]


def build_results_workbook(run, status="", search=""):
    """Only the Results rows (same filter as the Results screen).
    Normal (in-memory) workbook: no temp files, so Windows antivirus / file locks cannot stall it."""
    where, params = ["run_id = %s"], [run["run_id"]]
    if status:
        where.append("status = %s")
        params.append(status)
    if search:
        where.append("(packet_id ILIKE %s OR order_release_id ILIKE %s)")
        params += [f"%{search}%", f"%{search}%"]
    with db.cursor(commit=False) as cur:
        cur.execute(f"""SELECT status, packet_id, order_release_id, forward_sales, return_source, return_amount,
                               net_sales, forward_settlement, reverse_settlement, net_settlement, difference,
                               {MATCHED_AMOUNT_SQL} AS matched_amount, {MISMATCHED_AMOUNT_SQL} AS mismatched_amount
                        FROM recon_results WHERE {" AND ".join(where)}
                        ORDER BY status, packet_id, order_release_id""", params)
        rows = cur.fetchall()

    wb = Workbook()
    ws = wb.active
    ws.title = "Results"
    ws.append(RESULT_HEADERS)
    for c in ws[1]:
        c.font, c.fill = _HEAD_FONT, _HEAD_FILL
    money_idx = (3, 5, 6, 7, 8, 9, 10, 11, 12)
    for r in rows:
        ws.append([label(r["status"]), _num(r["packet_id"]), _num(r["order_release_id"]), _num(r["forward_sales"]),
                   None if r["return_source"] in (None, "NONE") else r["return_source"], _num(r["return_amount"]),
                   _num(r["net_sales"]), _num(r["forward_settlement"]), _num(r["reverse_settlement"]),
                   _num(r["net_settlement"]), _num(r["difference"]), _num(r["matched_amount"]),
                   _num(r["mismatched_amount"])])
    tot = {k: float(sum((r[k] or 0) for r in rows)) for k in
           ("net_sales", "net_settlement", "difference", "matched_amount", "mismatched_amount")}
    ws.append([f"TOTAL ({len(rows):,} rows)", None, None, None, None, None, tot["net_sales"], None, None,
               tot["net_settlement"], tot["difference"], tot["matched_amount"], tot["mismatched_amount"]])
    last = ws.max_row
    for c in ws[last]:
        c.font, c.fill = _BOLD, _TOTAL_FILL
    for row in ws.iter_rows(min_row=2, max_row=last):
        for i in money_idx:
            row[i].number_format = MONEY
    for i, w in enumerate([28, 20, 26, 16, 12, 16, 16, 18, 18, 16, 24, 18, 18], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(RESULT_HEADERS))}{max(last - 1, 2)}"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
